import csv
from decimal import ROUND_HALF_UP, Decimal
from itertools import groupby

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db.models import Sum
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone

from clientes.models import Cliente
from equipos.models import Equipo
from inventario.models import Producto
from rentas.models import DetalleRenta, Renta
from usuarios.decorators import cajero_o_admin_required
from usuarios.models import Usuario

from .models import DetalleVenta, Venta
from .ticket import generar_factura_pdf

IVA = Decimal('0.13')
CENTAVOS = Decimal('0.01')


@login_required
def inicio(request):
    hoy = timezone.localdate()

    ventas_hoy_qs = Venta.objects.filter(fecha__date=hoy)
    rentas_hoy_qs = Renta.objects.filter(fecha__date=hoy)

    # Ingresos totales de hoy (POS + Rentas)
    subtotal_ventas = sum((v.subtotal + v.iva) for v in ventas_hoy_qs) or Decimal('0.00')
    subtotal_rentas = sum(r.total for r in rentas_hoy_qs) or Decimal('0.00')
    ingresos_hoy = (subtotal_ventas + subtotal_rentas).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

    # Equipos en uso
    equipos_ocupados = DetalleRenta.objects.filter(estado='ACTIVA').values('equipo_id').distinct().count()
    total_equipos = Equipo.objects.count()
    equipos_activos_str = f"{equipos_ocupados} / {total_equipos}" if total_equipos else "0 / 0"

    # Total de ventas POS hoy
    ventas_hoy_count = ventas_hoy_qs.count()

    # Clientes activos únicos hoy (set union para deduplicación estricta)
    clientes_ventas = set(ventas_hoy_qs.exclude(cliente_id__isnull=True).values_list('cliente_id', flat=True))
    clientes_rentas = set(rentas_hoy_qs.exclude(cliente_id__isnull=True).values_list('cliente_id', flat=True))
    total_clientes_activos = len(clientes_ventas | clientes_rentas)

    # Últimas transacciones (Ventas y Rentas)
    ultimas_ventas_raw = list(
        Venta.objects.select_related('cliente', 'usuario')
        .prefetch_related('detalles__producto')
        .order_by('-fecha')[:6]
    )
    ultimas_rentas_raw = list(
        Renta.objects.select_related('cliente', 'usuario')
        .prefetch_related('detalles__equipo')
        .order_by('-fecha')[:6]
    )

    transacciones = []
    for v in ultimas_ventas_raw:
        items_nombres = [d.producto.nombre for d in v.detalles.all()[:2]]
        concepto = ", ".join(items_nombres) if items_nombres else "Venta de mostrador"
        if v.detalles.count() > 2:
            concepto += f" (+{v.detalles.count() - 2} más)"
        cliente_nombre = f"{v.cliente.nombre} {v.cliente.apellido}" if v.cliente else "Cliente General"
        transacciones.append({
            'id': f"TK-{v.venta_id}",
            'cliente': v.cliente,
            'cliente_nombre': cliente_nombre,
            'concepto': concepto,
            'usuario': v.usuario,
            'fecha_hora': v.fecha,
            'total': (v.subtotal + v.iva).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP),
            'estado': 'Pagado',
            'es_venta': True,
        })

    for r in ultimas_rentas_raw:
        det = r.detalles.first()
        equipo_nom = det.equipo.nombre if det and det.equipo else "Equipo"
        minutos = det.tiempo_de_renta if det else 60
        cliente_nombre = f"{r.cliente.nombre} {r.cliente.apellido}" if r.cliente else "Cliente General"
        transacciones.append({
            'id': f"RN-{r.renta_id}",
            'cliente': r.cliente,
            'cliente_nombre': cliente_nombre,
            'concepto': f"Renta {equipo_nom} ({minutos} min)",
            'usuario': r.usuario,
            'fecha_hora': r.fecha,
            'total': r.total,
            'estado': r.estado_de_pago.title(),
            'es_venta': False,
        })

    # Ordenar por fecha descendente y tomar las 5 más recientes
    transacciones.sort(key=lambda x: x['fecha_hora'], reverse=True)
    ultimas_transacciones = transacciones[:5]

    # Configuración de semana actual (Lunes a Domingo)
    inicio_semana = hoy - timezone.timedelta(days=hoy.weekday())  # Lunes
    fin_semana = inicio_semana + timezone.timedelta(days=6)       # Domingo

    # Gráfica 1: Ingresos de la Semana (Lunes a Domingo)
    dias_semana_es = ['Lunes', 'Martes', 'Miércoles', 'Jueves', 'Viernes', 'Sábado', 'Domingo']
    chart_dias_labels = []
    chart_ventas_data = []
    chart_rentas_data = []

    for i in range(7):
        dia = inicio_semana + timezone.timedelta(days=i)
        etiqueta = dias_semana_es[i]
        chart_dias_labels.append(etiqueta)

        v_sub = sum((v.subtotal + v.iva) for v in Venta.objects.filter(fecha__date=dia)) or Decimal('0.00')
        r_sub = sum(r.total for r in Renta.objects.filter(fecha__date=dia)) or Decimal('0.00')

        chart_ventas_data.append(float(v_sub))
        chart_rentas_data.append(float(r_sub))

    # Gráfica 2: Top 5 Productos Más Vendidos de la Semana (Lunes a Domingo)
    top_vendidos_qs = list(
        DetalleVenta.objects.filter(venta__fecha__date__range=(inicio_semana, fin_semana))
        .values('producto__nombre')
        .annotate(total_unidades=Sum('cantidad'))
        .order_by('-total_unidades')[:5]
    )
    chart_top_labels = []
    chart_top_data = []
    for item in top_vendidos_qs:
        if item['producto__nombre']:
            chart_top_labels.append(item['producto__nombre'])
            chart_top_data.append(int(item['total_unidades']))

    tiene_datos_top = bool(chart_top_labels)
    if not tiene_datos_top:
        chart_top_labels = ['Sin ventas esta semana']
        chart_top_data = [1]

    # Resumen semanal acumulado (Lunes a Domingo)
    total_semana_ventas = sum(chart_ventas_data)
    total_semana_rentas = sum(chart_rentas_data)
    total_semana = total_semana_ventas + total_semana_rentas

    # Ocupación de equipos
    equipos_libres = max(0, total_equipos - equipos_ocupados)
    porcentaje_ocupacion = int(round((equipos_ocupados / total_equipos * 100))) if total_equipos else 0

    import json

    return render(request, 'inicio.html', {
        'ingresos_hoy': ingresos_hoy,
        'equipos_activos': equipos_activos_str,
        'equipos_ocupados': equipos_ocupados,
        'equipos_libres': equipos_libres,
        'total_equipos': total_equipos,
        'porcentaje_ocupacion': porcentaje_ocupacion,
        'ventas_hoy': ventas_hoy_count,
        'total_clientes': total_clientes_activos,
        'ultimas_ventas': ultimas_transacciones,
        'inicio_semana': inicio_semana,
        'fin_semana': fin_semana,
        'chart_dias_labels': json.dumps(chart_dias_labels),
        'chart_ventas_data': json.dumps(chart_ventas_data),
        'chart_rentas_data': json.dumps(chart_rentas_data),
        'chart_top_labels': json.dumps(chart_top_labels),
        'chart_top_data': json.dumps(chart_top_data),
        'tiene_datos_top': tiene_datos_top,
        'total_semana': total_semana,
        'total_semana_ventas': total_semana_ventas,
        'total_semana_rentas': total_semana_rentas,
    })


def _obtener_carrito(request):
    carrito = request.session.get('carrito')
    if not isinstance(carrito, dict):
        return {}
    return carrito


def _resumen_carrito(carrito):
    items = []
    subtotal = Decimal('0.00')
    for producto_id, cantidad in carrito.items():
        producto = (
            Producto.objects.filter(pk=producto_id, activo=True)
            .select_related('categoria')
            .first()
        )
        if producto is None:
            continue
        linea = producto.precio_venta * cantidad
        subtotal += linea
        items.append({'producto': producto, 'cantidad': cantidad, 'subtotal': linea})
    iva = (subtotal * IVA).quantize(CENTAVOS, rounding=ROUND_HALF_UP)
    return items, subtotal, iva, subtotal + iva


def _respuesta_pdf(venta):
    respuesta = HttpResponse(generar_factura_pdf(venta), content_type='application/pdf')
    respuesta['Content-Disposition'] = f'inline; filename="factura_{venta.venta_id}.pdf"'
    return respuesta


@login_required
@cajero_o_admin_required
def pos(request):
    productos = (
        Producto.objects.filter(activo=True)
        .select_related('categoria')
        .order_by('categoria__nombre', 'nombre')
    )

    # Organizar por categorías
    categorias_dict = {}
    for p in productos:
        cat_nom = p.categoria.nombre.strip().title() if p.categoria else 'General'
        if cat_nom not in categorias_dict:
            categorias_dict[cat_nom] = []
        categorias_dict[cat_nom].append(p)

    categorias_resumen = [
        {'nombre': cat_nom, 'cantidad': len(prods)}
        for cat_nom, prods in categorias_dict.items()
    ]

    items, subtotal, iva, total = _resumen_carrito(_obtener_carrito(request))
    contexto = {
        'productos': productos,
        'categorias_dict': categorias_dict,
        'categorias_resumen': categorias_resumen,
        'items': items,
        'subtotal': subtotal,
        'iva': iva,
        'total': total,
        'clientes': Cliente.objects.order_by('nombre', 'apellido'),
        'iva_porcentaje': int(IVA * 100),
        'ultima_venta_id': request.session.get('ultima_venta_id'),
    }
    return render(request, 'pos/pos.html', contexto)


@login_required
@cajero_o_admin_required
def agregar_carrito(request):
    if request.method == 'POST':
        producto = get_object_or_404(Producto, pk=request.POST.get('producto_id'), activo=True)
        try:
            cantidad = int(request.POST.get('cantidad') or 1)
        except (TypeError, ValueError):
            cantidad = 1
        if cantidad < 1:
            cantidad = 1

        carrito = _obtener_carrito(request)
        clave = str(producto.producto_id)
        actual = carrito.get(clave, 0)
        if actual + cantidad > producto.cantidad:
            messages.error(
                request,
                f'Inventario insuficiente de "{producto.nombre}" '
                f'(disponible: {producto.cantidad}).',
            )
        else:
            carrito[clave] = actual + cantidad
            request.session['carrito'] = carrito
    return redirect('pos')


@login_required
@cajero_o_admin_required
def actualizar_cantidad_carrito(request):
    if request.method == 'POST':
        producto_id = str(request.POST.get('producto_id'))
        accion = request.POST.get('accion')
        producto = get_object_or_404(Producto, pk=producto_id, activo=True)
        carrito = _obtener_carrito(request)
        actual = carrito.get(producto_id, 0)

        if accion == 'incrementar':
            if actual + 1 > producto.cantidad:
                messages.error(
                    request,
                    f'Inventario insuficiente de "{producto.nombre}" (disponible: {producto.cantidad}).',
                )
            else:
                carrito[producto_id] = actual + 1
        elif accion == 'decrementar':
            if actual - 1 <= 0:
                carrito.pop(producto_id, None)
            else:
                carrito[producto_id] = actual - 1
        request.session['carrito'] = carrito
    return redirect('pos')


@login_required
@cajero_o_admin_required
def quitar_carrito(request):
    if request.method == 'POST':
        carrito = _obtener_carrito(request)
        carrito.pop(str(request.POST.get('producto_id')), None)
        request.session['carrito'] = carrito
    return redirect('pos')


@login_required
@cajero_o_admin_required
def limpiar_carrito(request):
    if request.method == 'POST':
        request.session['carrito'] = {}
        messages.info(request, 'El carrito ha sido vaciado.')
    return redirect('pos')


@login_required
@cajero_o_admin_required
def procesar_venta(request):
    if request.method != 'POST':
        return redirect('pos')

    items, subtotal, iva, _total = _resumen_carrito(_obtener_carrito(request))
    if not items:
        messages.error(request, 'El carrito está vacío. Agrega productos antes de procesar el pago.')
        return redirect('pos')

    cliente = get_object_or_404(Cliente, pk=request.POST.get('cliente_id'))

    venta = Venta.objects.create(
        cliente=cliente,
        usuario=request.user,
        fecha=timezone.now(),
        subtotal=subtotal,
        iva=iva,
    )

    for item in items:
        producto = item['producto']
        DetalleVenta.objects.create(
            venta=venta,
            producto=producto,
            precio_unitario=producto.precio_venta,
            cantidad=item['cantidad'],
            subtotal=item['subtotal'],
        )
        producto.cantidad = max(producto.cantidad - item['cantidad'], 0)
        producto.save(update_fields=['cantidad'])

    request.session['carrito'] = {}
    request.session['ultima_venta_id'] = venta.venta_id
    messages.success(request, f'Venta #{venta.venta_id} registrada correctamente.')

    return _respuesta_pdf(venta)


@login_required
@cajero_o_admin_required
def ticket_pdf(request, venta_id=None):
    venta = None
    if venta_id is not None:
        venta = Venta.objects.filter(pk=venta_id).first()
    else:
        ultima = request.session.get('ultima_venta_id')
        if ultima:
            venta = Venta.objects.filter(pk=ultima).first()

    if venta is None:
        messages.error(request, 'No hay ninguna venta registrada para generar el ticket.')
        return redirect('pos')

    return _respuesta_pdf(venta)


@login_required
@cajero_o_admin_required
def reportes(request):
    fecha_inicio_str = request.GET.get('fecha_inicio')
    fecha_fin_str = request.GET.get('fecha_fin')
    tipo_filtro = request.GET.get('tipo', 'todos')
    usuario_id_str = request.GET.get('usuario')

    ventas_qs = Venta.objects.select_related('cliente', 'usuario').prefetch_related('detalles__producto').order_by('-fecha')
    rentas_qs = Renta.objects.select_related('cliente', 'usuario').prefetch_related('detalles__equipo').order_by('-fecha')

    if fecha_inicio_str:
        ventas_qs = ventas_qs.filter(fecha__date__gte=fecha_inicio_str)
        rentas_qs = rentas_qs.filter(fecha__date__gte=fecha_inicio_str)
    if fecha_fin_str:
        ventas_qs = ventas_qs.filter(fecha__date__lte=fecha_fin_str)
        rentas_qs = rentas_qs.filter(fecha__date__lte=fecha_fin_str)

    if usuario_id_str and usuario_id_str.isdigit():
        uid = int(usuario_id_str)
        ventas_qs = ventas_qs.filter(usuario_id=uid)
        rentas_qs = rentas_qs.filter(usuario_id=uid)

    transacciones = []
    total_ventas = Decimal('0.00')
    total_rentas = Decimal('0.00')

    if tipo_filtro in ['todos', 'ventas']:
        for v in ventas_qs:
            items_nombres = [f"{d.cantidad}x {d.producto.nombre}" for d in v.detalles.all()]
            articulos = ", ".join(items_nombres) if items_nombres else "Venta de mostrador"
            monto_v = (v.subtotal + v.iva).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
            total_ventas += monto_v
            transacciones.append({
                'ticket': f"V-{v.venta_id}",
                'fecha': v.fecha,
                'usuario': v.usuario.username if v.usuario else "Sistema",
                'cliente': f"{v.cliente.nombre} {v.cliente.apellido}" if v.cliente else "Cliente General",
                'concepto': articulos,
                'total': monto_v,
                'tipo': 'Venta POS',
                'es_venta': True,
                'venta_id': v.venta_id,
            })

    if tipo_filtro in ['todos', 'rentas']:
        for r in rentas_qs:
            det = r.detalles.first()
            equipo_nom = det.equipo.nombre if det and det.equipo else "Equipo"
            minutos = (det.tiempo_de_renta + (det.tiempo_agregado or 0)) if det else 60
            monto_r = r.total.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
            total_rentas += monto_r
            transacciones.append({
                'ticket': f"R-{r.renta_id}",
                'fecha': r.fecha,
                'usuario': r.usuario.username if r.usuario else "Sistema",
                'cliente': f"{r.cliente.nombre} {r.cliente.apellido}" if r.cliente else "Cliente General",
                'concepto': f"Renta {equipo_nom} ({minutos} min)",
                'total': monto_r,
                'tipo': 'Renta de Equipo',
                'es_venta': False,
                'venta_id': None,
            })

    transacciones.sort(key=lambda x: x['fecha'], reverse=True)
    total_global = total_ventas + total_rentas
    num_transacciones = len(transacciones)
    ticket_promedio = (total_global / Decimal(num_transacciones)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP) if num_transacciones else Decimal('0.00')

    pct_ventas = int((total_ventas / total_global * 100)) if total_global > 0 else 0
    pct_rentas = 100 - pct_ventas if total_global > 0 else 0

    usuarios = Usuario.objects.filter(activo=True).order_by('username')

    return render(request, 'reportes.html', {
        'transacciones': transacciones,
        'total_recaudado_global': total_global,
        'total_ventas': total_ventas,
        'total_rentas': total_rentas,
        'num_transacciones': num_transacciones,
        'ticket_promedio': ticket_promedio,
        'pct_ventas': pct_ventas,
        'pct_rentas': pct_rentas,
        'fecha_inicio': fecha_inicio_str or '',
        'fecha_fin': fecha_fin_str or '',
        'tipo_filtro': tipo_filtro,
        'usuario_id_seleccionado': int(usuario_id_str) if (usuario_id_str and usuario_id_str.isdigit()) else '',
        'usuarios': usuarios,
        'ahora': timezone.now(),
    })


@login_required
@cajero_o_admin_required
def reportes_exportar_csv(request):
    fecha_inicio_str = request.GET.get('fecha_inicio')
    fecha_fin_str = request.GET.get('fecha_fin')
    tipo_filtro = request.GET.get('tipo', 'todos')
    usuario_id_str = request.GET.get('usuario')

    ventas_qs = Venta.objects.select_related('cliente', 'usuario').prefetch_related('detalles__producto').order_by('-fecha')
    rentas_qs = Renta.objects.select_related('cliente', 'usuario').prefetch_related('detalles__equipo').order_by('-fecha')

    if fecha_inicio_str:
        ventas_qs = ventas_qs.filter(fecha__date__gte=fecha_inicio_str)
        rentas_qs = rentas_qs.filter(fecha__date__gte=fecha_inicio_str)
    if fecha_fin_str:
        ventas_qs = ventas_qs.filter(fecha__date__lte=fecha_fin_str)
        rentas_qs = rentas_qs.filter(fecha__date__lte=fecha_fin_str)

    if usuario_id_str and usuario_id_str.isdigit():
        uid = int(usuario_id_str)
        ventas_qs = ventas_qs.filter(usuario_id=uid)
        rentas_qs = rentas_qs.filter(usuario_id=uid)

    filas = []
    if tipo_filtro in ['todos', 'ventas']:
        for v in ventas_qs:
            items_nombres = [f"{d.cantidad}x {d.producto.nombre}" for d in v.detalles.all()]
            articulos = ", ".join(items_nombres) if items_nombres else "Venta de mostrador"
            filas.append({
                'fecha_obj': v.fecha,
                'ticket': f"V-{v.venta_id}",
                'tipo': "Venta POS",
                'fecha': v.fecha.strftime('%d/%m/%Y %I:%M %p'),
                'cajero': v.usuario.username if v.usuario else "Sistema",
                'cliente': f"{v.cliente.nombre} {v.cliente.apellido}" if v.cliente else "Cliente General",
                'concepto': articulos,
                'total': f"${(v.subtotal + v.iva).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)}",
            })

    if tipo_filtro in ['todos', 'rentas']:
        for r in rentas_qs:
            det = r.detalles.first()
            equipo_nom = det.equipo.nombre if det and det.equipo else "Equipo"
            minutos = (det.tiempo_de_renta + (det.tiempo_agregado or 0)) if det else 60
            filas.append({
                'fecha_obj': r.fecha,
                'ticket': f"R-{r.renta_id}",
                'tipo': "Renta de Equipo",
                'fecha': r.fecha.strftime('%d/%m/%Y %I:%M %p'),
                'cajero': r.usuario.username if r.usuario else "Sistema",
                'cliente': f"{r.cliente.nombre} {r.cliente.apellido}" if r.cliente else "Cliente General",
                'concepto': f"Renta {equipo_nom} ({minutos} min)",
                'total': f"${r.total.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)}",
            })

    filas.sort(key=lambda x: x['fecha_obj'], reverse=True)

    response = HttpResponse(content_type='text/csv; charset=utf-8')
    response['Content-Disposition'] = f'attachment; filename="reporte_ingresos_{timezone.localdate()}.csv"'
    response.write('\ufeff')  # UTF-8 BOM para Excel en español

    writer = csv.writer(response)
    writer.writerow(['Ticket', 'Tipo de Ingreso', 'Fecha y Hora', 'Cajero / Atendido Por', 'Cliente', 'Detalle / Concepto', 'Total'])
    for f in filas:
        writer.writerow([f['ticket'], f['tipo'], f['fecha'], f['cajero'], f['cliente'], f['concepto'], f['total']])

    return response


@login_required
@cajero_o_admin_required
def reportes_exportar_excel(request):
    fecha_inicio_str = request.GET.get('fecha_inicio')
    fecha_fin_str = request.GET.get('fecha_fin')
    tipo_filtro = request.GET.get('tipo', 'todos')
    usuario_id_str = request.GET.get('usuario')

    ventas_qs = Venta.objects.select_related('cliente', 'usuario').prefetch_related('detalles__producto').order_by('-fecha')
    rentas_qs = Renta.objects.select_related('cliente', 'usuario').prefetch_related('detalles__equipo').order_by('-fecha')

    if fecha_inicio_str:
        ventas_qs = ventas_qs.filter(fecha__date__gte=fecha_inicio_str)
        rentas_qs = rentas_qs.filter(fecha__date__gte=fecha_inicio_str)
    if fecha_fin_str:
        ventas_qs = ventas_qs.filter(fecha__date__lte=fecha_fin_str)
        rentas_qs = rentas_qs.filter(fecha__date__lte=fecha_fin_str)

    if usuario_id_str and usuario_id_str.isdigit():
        uid = int(usuario_id_str)
        ventas_qs = ventas_qs.filter(usuario_id=uid)
        rentas_qs = rentas_qs.filter(usuario_id=uid)

    filas = []
    if tipo_filtro in ['todos', 'ventas']:
        for v in ventas_qs:
            items_nombres = [f"{d.cantidad}x {d.producto.nombre}" for d in v.detalles.all()]
            articulos = ", ".join(items_nombres) if items_nombres else "Venta de mostrador"
            filas.append({
                'fecha_obj': v.fecha,
                'ticket': f"V-{v.venta_id}",
                'tipo': "Venta POS",
                'fecha': v.fecha.strftime('%d/%m/%Y %I:%M %p'),
                'cajero': v.usuario.username if v.usuario else "Sistema",
                'cliente': f"{v.cliente.nombre} {v.cliente.apellido}" if v.cliente else "Cliente General",
                'concepto': articulos,
                'total': float((v.subtotal + v.iva).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)),
            })

    if tipo_filtro in ['todos', 'rentas']:
        for r in rentas_qs:
            det = r.detalles.first()
            equipo_nom = det.equipo.nombre if det and det.equipo else "Equipo"
            minutos = (det.tiempo_de_renta + (det.tiempo_agregado or 0)) if det else 60
            filas.append({
                'fecha_obj': r.fecha,
                'ticket': f"R-{r.renta_id}",
                'tipo': "Renta de Equipo",
                'fecha': r.fecha.strftime('%d/%m/%Y %I:%M %p'),
                'cajero': r.usuario.username if r.usuario else "Sistema",
                'cliente': f"{r.cliente.nombre} {r.cliente.apellido}" if r.cliente else "Cliente General",
                'concepto': f"Renta {equipo_nom} ({minutos} min)",
                'total': float(r.total.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)),
            })

    filas.sort(key=lambda x: x['fecha_obj'], reverse=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Reporte de Ingresos"

    headers = ['N° Ticket', 'Tipo de Ingreso', 'Fecha y Hora', 'Cajero / Atendido Por', 'Cliente', 'Detalle / Concepto', 'Total ($)']
    ws.append(headers)

    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='2C1E16', end_color='2C1E16', fill_type='solid')
    header_alignment = Alignment(horizontal='center', vertical='center')

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    for row_idx, f in enumerate(filas, start=2):
        ws.cell(row=row_idx, column=1, value=f['ticket'])
        ws.cell(row=row_idx, column=2, value=f['tipo'])
        ws.cell(row=row_idx, column=3, value=f['fecha'])
        ws.cell(row=row_idx, column=4, value=f['cajero'])
        ws.cell(row=row_idx, column=5, value=f['cliente'])
        ws.cell(row=row_idx, column=6, value=f['concepto'])
        total_cell = ws.cell(row=row_idx, column=7, value=f['total'])
        total_cell.number_format = '$#,##0.00'

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="reporte_ingresos_{timezone.localdate()}.xlsx"'
    wb.save(response)
    return response

